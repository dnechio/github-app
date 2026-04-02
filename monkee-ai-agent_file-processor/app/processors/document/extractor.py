from app.processors.base import BaseProcessor
from langchain_text_splitters import RecursiveCharacterTextSplitter
from loguru import logger


class DocumentProcessor(BaseProcessor):
    """Extracts text from DOCX, TXT, XLSX, HTML, MD files."""

    def extract_text(self) -> str:
        from app.storage.s3 import download_file_to_tmp
        tmp_path = download_file_to_tmp(self.file_id)
        ext = tmp_path.suffix.lower()

        if ext in (".docx", ".doc"):
            return self._extract_docx(tmp_path)
        elif ext in (".xlsx", ".xls"):
            return self._extract_xlsx(tmp_path)
        elif ext in (".txt", ".md"):
            return tmp_path.read_text(encoding="utf-8", errors="ignore")
        elif ext == ".html":
            return self._extract_html(tmp_path)
        else:
            logger.warning(f"Unknown doc extension: {ext}")
            return tmp_path.read_text(encoding="utf-8", errors="ignore")

    def _extract_docx(self, path) -> str:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def _extract_xlsx(self, path) -> str:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                rows.append("\t".join(str(c) for c in row if c is not None))
        return "\n".join(rows)

    def _extract_html(self, path) -> str:
        import re
        html = path.read_text(encoding="utf-8", errors="ignore")
        return re.sub(r"<[^>]+>", " ", html).strip()

    def chunk_text(self, text: str) -> list[str]:
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        return splitter.split_text(text)
